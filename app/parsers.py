from __future__ import annotations

import io
import re
import zipfile
from pathlib import Path

from docx import Document
from pypdf import PdfReader

RESUME_EXTS = {".pdf", ".docx", ".doc", ".txt", ".md"}
MAX_ZIP_FILES = 5000
MAX_FILE_BYTES = 20 * 1024 * 1024

EMAIL_RE = re.compile(r"[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Za-z]{2,}")
PHONE_RE = re.compile(
    r"(?:\+?\d{1,3}[\s.-]?)?(?:\(?\d{2,4}\)?[\s.-]?)?\d{3,4}[\s.-]?\d{3,4}"
)
SKIP_NAME = re.compile(
    r"^(resume|curriculum vitae|cv|profile|contact|objective|summary)$",
    re.I,
)


def is_resume_zip_entry(info: zipfile.ZipInfo) -> bool:
    if info.is_dir() or info.file_size > MAX_FILE_BYTES:
        return False
    path = Path(info.filename)
    if ".." in path.parts or path.name.startswith("."):
        return False
    if "__MACOSX" in path.parts:
        return False
    return path.suffix.lower() in RESUME_EXTS


def count_resumes(filename: str, data: bytes | None = None, path: Path | None = None) -> int:
    suffix = Path(filename).suffix.lower()
    if suffix == ".zip":
        archive = zipfile.ZipFile(path) if path else zipfile.ZipFile(io.BytesIO(data or b""))
        with archive:
            n = sum(1 for info in archive.infolist() if is_resume_zip_entry(info))
        return min(n, MAX_ZIP_FILES)
    if suffix in RESUME_EXTS:
        return 1
    return 0


def iter_resumes(filename: str, data: bytes | None = None, path: Path | None = None):
    suffix = Path(filename).suffix.lower()
    if suffix == ".zip":
        archive = zipfile.ZipFile(path) if path else zipfile.ZipFile(io.BytesIO(data or b""))
        yielded = 0
        with archive:
            for info in archive.infolist():
                if not is_resume_zip_entry(info):
                    continue
                yield Path(info.filename).name, archive.read(info)
                yielded += 1
                if yielded >= MAX_ZIP_FILES:
                    break
        if yielded == 0:
            raise ValueError("Zip had no PDF, DOCX, or TXT resumes inside.")
        return
    if suffix not in RESUME_EXTS:
        raise ValueError("Use PDF, DOCX, TXT, or a zip of those files.")
    payload = path.read_bytes() if path else (data or b"")
    yield Path(filename).name, payload


def expand_uploads(filename: str, data: bytes) -> list[tuple[str, bytes]]:
    return list(iter_resumes(filename, data=data))


_ocr_engine = None


def parse_file(filename: str, data: bytes) -> str:
    suffix = Path(filename).suffix.lower()
    if suffix == ".pdf":
        return _parse_pdf(data)
    if suffix in {".docx", ".doc"}:
        return _parse_docx(data)
    if suffix in {".txt", ".md"}:
        return data.decode("utf-8", errors="replace")
    raise ValueError(f"Unsupported file type: {suffix or 'unknown'}")


def _parse_pdf(data: bytes) -> str:
    for extractor in (_pdf_pymupdf, _pdf_pdfminer, _pdf_pypdf):
        try:
            text = extractor(data)
        except Exception:
            continue
        if len(text) >= 40:
            return text
    try:
        text = _pdf_ocr(data)
    except Exception:
        text = ""
    return text


def _pdf_pymupdf(data: bytes) -> str:
    import pymupdf

    doc = pymupdf.open(stream=data, filetype="pdf")
    try:
        if doc.is_encrypted:
            doc.authenticate("")
        pages = [page.get_text("text", sort=True) or "" for page in doc]
    finally:
        doc.close()
    return _normalize("\n".join(pages))


def _pdf_pdfminer(data: bytes) -> str:
    from pdfminer.high_level import extract_text

    return _normalize(extract_text(io.BytesIO(data)) or "")


def _pdf_pypdf(data: bytes) -> str:
    reader = PdfReader(io.BytesIO(data))
    if reader.is_encrypted:
        try:
            reader.decrypt("")
        except Exception:
            pass
    pages = [page.extract_text() or "" for page in reader.pages]
    return _normalize("\n".join(pages))


def _pdf_ocr(data: bytes) -> str:
    import pymupdf

    engine = _get_ocr()
    doc = pymupdf.open(stream=data, filetype="pdf")
    lines: list[str] = []
    try:
        if doc.is_encrypted:
            doc.authenticate("")
        for index, page in enumerate(doc):
            if index >= 8:
                break
            pix = page.get_pixmap(matrix=pymupdf.Matrix(2, 2), alpha=False)
            result, _ = engine(pix.tobytes("png"))
            if not result:
                continue
            for item in result:
                if isinstance(item, (list, tuple)) and len(item) >= 2:
                    lines.append(str(item[1]))
    finally:
        doc.close()
    return _normalize("\n".join(lines))


def _get_ocr():
    global _ocr_engine
    if _ocr_engine is None:
        from rapidocr_onnxruntime import RapidOCR

        _ocr_engine = RapidOCR()
    return _ocr_engine


def _parse_docx(data: bytes) -> str:
    doc = Document(io.BytesIO(data))
    parts = [p.text for p in doc.paragraphs if p.text and p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return _normalize("\n".join(parts))


def _normalize(text: str) -> str:
    text = text.replace("\x00", " ")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def extract_profile(text: str, filename: str) -> dict:
    emails = EMAIL_RE.findall(text)
    phones = [
        p.strip()
        for p in PHONE_RE.findall(text[:2000])
        if len(re.sub(r"\D", "", p)) >= 10
    ]
    return {
        "name": _guess_name(text, filename),
        "email": emails[0] if emails else "",
        "phone": phones[0] if phones else "",
    }


def _guess_name(text: str, filename: str) -> str:
    head = text[:1200].splitlines()
    for raw in head[:12]:
        line = raw.strip(" |-•\t")
        if not line or len(line) > 60:
            continue
        if EMAIL_RE.search(line) or PHONE_RE.search(line):
            continue
        if SKIP_NAME.match(line):
            continue
        if any(ch.isdigit() for ch in line):
            continue
        words = [w for w in re.split(r"\s+", line) if w]
        if 2 <= len(words) <= 5 and all(w[0].isalpha() for w in words if w):
            return " ".join(words)
        if len(words) == 1 and words[0][0].isupper() and len(words[0]) > 2:
            return words[0]
    stem = Path(filename).stem.replace("_", " ").replace("-", " ").strip()
    return stem.title() if stem else "Unknown candidate"
