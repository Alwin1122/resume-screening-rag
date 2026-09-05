"""Export REPORT.md to Word and copy both files into the user's Downloads folder."""

from __future__ import annotations

import re
import shutil
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
SRC = ROOT / "REPORT.md"
DOWNLOADS = Path.home() / "Downloads"
STEM = "Resume_Screening_RAG_Project_Report"


def set_run_font(run, name: str = "Calibri", size: int = 11, bold: bool = False, italic: bool = False):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = RGBColor(0x1A, 0x1A, 0x1A)


def add_formatted(paragraph, text: str, size: int = 11):
    parts = re.split(r"(\*\*[^*]+\*\*|`[^`]+`)", text)
    for part in parts:
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            set_run_font(run, size=size, bold=True)
        elif part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            set_run_font(run, name="Consolas", size=size)
        else:
            run = paragraph.add_run(part)
            set_run_font(run, size=size)


def add_table(doc: Document, rows: list[list[str]]):
    if not rows:
        return
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    table.style = "Table Grid"
    for i, row in enumerate(rows):
        for j, cell in enumerate(row):
            table.cell(i, j).text = ""
            p = table.cell(i, j).paragraphs[0]
            add_formatted(p, cell.strip(), size=10)
            if i == 0:
                for run in p.runs:
                    run.bold = True
    doc.add_paragraph()


def parse_tables_and_blocks(md: str) -> list[tuple[str, object]]:
    lines = md.replace("\r\n", "\n").split("\n")
    blocks: list[tuple[str, object]] = []
    i = 0
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        if stripped.startswith("|") and i + 1 < len(lines) and re.match(r"^\|?\s*-+", lines[i + 1].strip()):
            table: list[list[str]] = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                row = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                if not re.match(r"^:?-+:?$", row[0].replace(" ", "")):
                    table.append(row)
                i += 1
            blocks.append(("table", table))
            continue

        if stripped.startswith("```"):
            lang = stripped[3:].strip()
            i += 1
            code: list[str] = []
            while i < len(lines) and not lines[i].strip().startswith("```"):
                code.append(lines[i])
                i += 1
            i += 1
            blocks.append(("code", "\n".join(code)))
            continue

        if stripped == "---":
            blocks.append(("hr", None))
            i += 1
            continue

        if stripped.startswith("# "):
            blocks.append(("h1", stripped[2:].strip()))
        elif stripped.startswith("## "):
            blocks.append(("h2", stripped[3:].strip()))
        elif stripped.startswith("### "):
            blocks.append(("h3", stripped[4:].strip()))
        elif re.match(r"^\d+\.\s+", stripped):
            blocks.append(("ol", re.sub(r"^\d+\.\s+", "", stripped)))
        elif stripped.startswith("- "):
            blocks.append(("ul", stripped[2:].strip()))
        elif stripped:
            blocks.append(("p", stripped))
        else:
            blocks.append(("blank", None))
        i += 1
    return blocks


def build_docx(md: str, dest: Path) -> None:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.9)
    section.bottom_margin = Inches(0.9)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)

    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)
    style.paragraph_format.space_after = Pt(6)
    style.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE

    first_h1 = True
    for kind, value in parse_tables_and_blocks(md):
        if kind == "h1":
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(value)
            set_run_font(run, name="Calibri", size=22, bold=True)
            p.paragraph_format.space_after = Pt(6)
            if first_h1:
                first_h1 = False
        elif kind == "h2":
            p = doc.add_heading(value, level=1)
            for run in p.runs:
                run.font.color.rgb = RGBColor(0x1B, 0x3A, 0x4B)
        elif kind == "h3":
            p = doc.add_heading(value, level=2)
            for run in p.runs:
                run.font.color.rgb = RGBColor(0x2C, 0x5F, 0x6E)
        elif kind == "p":
            p = doc.add_paragraph()
            add_formatted(p, value)
        elif kind == "ol":
            p = doc.add_paragraph(style="List Number")
            add_formatted(p, value)
        elif kind == "ul":
            p = doc.add_paragraph(style="List Bullet")
            add_formatted(p, value)
        elif kind == "table":
            add_table(doc, value)
        elif kind == "code":
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.25)
            run = p.add_run(value)
            set_run_font(run, name="Consolas", size=9)
        elif kind == "hr":
            continue

    dest.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(dest))


def main() -> None:
    if not SRC.exists():
        raise SystemExit(f"Missing {SRC}")
    md = SRC.read_text(encoding="utf-8")

    project_docx = ROOT / "REPORT.docx"
    downloads_docx = DOWNLOADS / f"{STEM}.docx"
    downloads_md = DOWNLOADS / f"{STEM}.md"

    build_docx(md, project_docx)
    shutil.copy2(project_docx, downloads_docx)
    shutil.copy2(SRC, downloads_md)

    print(downloads_docx)
    print(downloads_md)
    print(project_docx)


if __name__ == "__main__":
    main()
